"""生成真实解析样本（data/samples/，可重复运行）。

样本集对应方案第 13/16 节与 V01-V24 附录 A：
- normal.pdf    正常文字层 PDF（要素齐全）
- scanned.pdf   扫描件（图片页、无文字层，需 OCR）
- blur_name.pdf 姓名区域高斯模糊（V07：OCR 误识别不出姓名，其余可读）
- report.docx   DOCX（段落 + 表格）
- photo.png     图片材料（实物赛道照片，需 OCR）
- partial.pdf   2 页 PDF：第 1 页文字层、第 2 页图片页（V08 覆盖缺口）
- corrupt.pdf   损坏文件（技术原因，非资格拒绝）
- disguised.pdf 类型伪装（DOCX 内容谎称 .pdf）
- legacy.doc    旧版 OLE2 文档（不支持）
- encrypted.pdf 加密 PDF（AES-256，无密码无法读取）

用法：python scripts/make_samples.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Test fixtures live beside the repository's data directory. Keep this path
# independent from FORMFLOW_STORAGE_DIR, which the test suite overrides with a
# temporary private store before importing the application settings.
OUT = Path(__file__).resolve().parent.parent / "data" / "samples"

NAME = "张三"
TERM = "2026 秋季第 1 期"
CONTENT = "活动内容：参加了本期团课，完成了小组实验与讨论，承担了数据记录工作。"
GAIN = "个人收获：提升了实验操作与团队协作能力，认识了更多同学。"


def _write_pdf_text(path: Path) -> None:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    # china-s：PyMuPDF 内置简体中文字体（helv 不含 CJK，中文会变点）
    kw = {"fontname": "china-s"}
    page.insert_text((72, 72), f"{TERM}团课报告", fontsize=18, **kw)
    page.insert_text((72, 120), f"姓名：{NAME}", fontsize=14, **kw)
    page.insert_text((72, 150), f"期次：{TERM}", fontsize=14, **kw)
    page.insert_text((72, 180), CONTENT, fontsize=12, **kw)
    page.insert_text((72, 210), GAIN, fontsize=12, **kw)
    doc.save(str(path))
    doc.close()


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)

    # 1. 正常 PDF（有文字层）
    _write_pdf_text(OUT / "normal.pdf")

    # 2. 扫描件 PDF：先渲染成图，再贴进新 PDF（无文字层，必须 OCR）
    import fitz

    tmp = OUT / "_tmp_text.pdf"
    _write_pdf_text(tmp)
    src = fitz.open(str(tmp))
    pix = src[0].get_pixmap(dpi=150)
    img = OUT / "_tmp_page.png"
    pix.save(str(img))
    src.close()
    scanned = fitz.open()
    sp = scanned.new_page(width=pix.width, height=pix.height)
    sp.insert_image(fitz.Rect(0, 0, pix.width, pix.height), filename=str(img))
    scanned.save(str(OUT / "scanned.pdf"))
    scanned.close()
    tmp.unlink()
    img.unlink()

    # 3. DOCX（段落 + 表格）
    import docx

    d = docx.Document()
    d.add_heading(f"{TERM}团课报告", 0)
    d.add_paragraph(f"姓名：{NAME}")
    d.add_paragraph(f"期次：{TERM}")
    d.add_paragraph(CONTENT)
    d.add_paragraph(GAIN)
    tbl = d.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "姓名"
    tbl.rows[0].cells[1].text = NAME
    tbl.rows[1].cells[0].text = "期次"
    tbl.rows[1].cells[1].text = TERM
    d.save(str(OUT / "report.docx"))

    # 4. 照片 PNG：渲染文本为图片
    tmp2 = OUT / "_tmp_text2.pdf"
    _write_pdf_text(tmp2)
    d2 = fitz.open(str(tmp2))
    d2[0].get_pixmap(dpi=150).save(str(OUT / "photo.png"))
    d2.close()
    tmp2.unlink()

    # 4b. 姓名区域模糊的 PDF（V07）：其余要素可读，姓名行被高斯模糊，
    #     OCR 误识别辨不出姓名——模拟“匿名字缺失”的真实场景。
    tmp3 = OUT / "_tmp_text3.pdf"
    _write_pdf_text(tmp3)
    _d3 = fitz.open(str(tmp3))
    _pix3 = _d3[0].get_pixmap(dpi=150)
    _d3.close()
    tmp3.unlink()
    from PIL import Image, ImageFilter

    _im = Image.frombytes("RGB", (_pix3.width, _pix3.height), _pix3.samples).copy()
    _scale = 150 / 72
    _box = (int(70 * _scale), int(106 * _scale),
            int(340 * _scale), int(140 * _scale))
    _im.paste(_im.crop(_box).filter(ImageFilter.GaussianBlur(6)), _box)
    _blurred = OUT / "_tmp_blur.png"
    _im.save(str(_blurred))
    _bd = fitz.open()
    _bd.new_page(width=_pix3.width, height=_pix3.height).insert_image(
        fitz.Rect(0, 0, _pix3.width, _pix3.height), filename=str(_blurred))
    _bd.save(str(OUT / "blur_name.pdf"))
    _bd.close()
    _blurred.unlink()

    # 4c. partial.pdf（V08）：2 页——第 1 页文字层要素齐全（但无“个人收获”），
    #     第 2 页为图片页（无文字层，承载“个人收获”）。
    #     关键页未读时不得输出“确认缺失”，覆盖缺口必须可见。
    _one = fitz.open()
    _one.new_page().insert_text((72, 72), f"{TERM}团课报告", fontsize=18, fontname="china-s")
    _one[0].insert_text((72, 120), f"姓名：{NAME}", fontsize=14, fontname="china-s")
    _one[0].insert_text((72, 150), f"期次：{TERM}", fontsize=14, fontname="china-s")
    _one[0].insert_text((72, 180), CONTENT, fontsize=12, fontname="china-s")
    _one.save(str(OUT / "_tmp_p1.pdf"))
    _one.close()
    _two = fitz.open()
    _two.new_page().insert_text((72, 72), "个人收获", fontsize=14, fontname="china-s")
    _two[0].insert_text((72, 100), GAIN, fontsize=12, fontname="china-s")
    _two.save(str(OUT / "_tmp_p2.pdf"))
    _two.close()
    _pix2 = fitz.open(str(OUT / "_tmp_p2.pdf"))[0].get_pixmap(dpi=150)
    _img2 = OUT / "_tmp_p2.png"
    _pix2.save(str(_img2))
    _merged = fitz.open()
    _src1 = fitz.open(str(OUT / "_tmp_p1.pdf"))
    _merged.insert_pdf(_src1)
    _src1.close()
    _merged.new_page(width=_pix2.width, height=_pix2.height).insert_image(
        fitz.Rect(0, 0, _pix2.width, _pix2.height), filename=str(_img2))
    _merged.save(str(OUT / "partial.pdf"))
    _merged.close()
    _img2.unlink()
    (OUT / "_tmp_p1.pdf").unlink()
    (OUT / "_tmp_p2.pdf").unlink()

    # 5. 损坏 PDF：合法头 + 垃圾
    (OUT / "corrupt.pdf").write_bytes(b"%PDF-1.4" + bytes(range(256)) * 8)

    # 6. 类型伪装：DOCX 内容、.pdf 扩展名
    shutil.copyfile(OUT / "report.docx", OUT / "disguised.pdf")

    # 7. 旧版 OLE2 文档
    (OUT / "legacy.doc").write_bytes(bytes([0xD0, 0xCF, 0x11, 0xE0]) + b"legacy junk")
    # 8. 加密 PDF：AES-256，无密码读不出内容（V11 技术原因，非资格拒绝）
    enc = fitz.open()
    ep = enc.new_page()
    ep.insert_text((72, 72), "加密内容", fontsize=14, fontname="china-s")
    enc.save(str(OUT / "encrypted.pdf"), encryption=fitz.PDF_ENCRYPT_AES_256,
             owner_pw="owner-secret", user_pw="user-secret")
    enc.close()

    print("samples written to", OUT)
    for f in sorted(OUT.iterdir()):
        print(" -", f.name, f.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
