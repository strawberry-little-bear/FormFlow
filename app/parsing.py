"""文档解析：提取文本用于语义检查（第 16 节边界）。

安全与资源边界：
- 只读解析，不执行宏、脚本或上传代码；损坏/加密/类型伪装返回 ParseError，
  调用方按技术原因处理，不当成资格拒绝。
- OCR 是尽力而为：扫描件/图片页按需 OCR（PyMuPDF 渲染 + RapidOCR 本地识别）。
  识别失败、引擎缺失或超出预算的页计入未读页，语义结论必须 unknown（V07），
  不得把 OCR 缺失写成"未填写"。
- 资源限制：单文件 OCR 页数上限与整体时限，超限部分记未读，不无限消耗。
"""
from __future__ import annotations

import time
from dataclasses import dataclass, field
from pathlib import Path

from .config import settings


class ParseError(Exception):
    """解析失败（损坏/加密/格式不支持）。调用方按技术原因处理。"""


@dataclass
class ParseResult:
    pages: list[str] = field(default_factory=list)
    method: str = "text"            # text / ocr / mixed / visual / none
    unread_pages: int = 0           # 需要 OCR 但失败/超限的页
    failed_pages: int = 0
    ocr_pages: int = 0              # 经 OCR 成功提取文字的页数
    detected: str = ""              # 按内容检测到的真实类型：pdf/docx/image/text
    error: str = ""

    @property
    def total_pages(self) -> int:
        return len(self.pages) + self.unread_pages + self.failed_pages

    @property
    def text_pages(self) -> int:
        return sum(1 for p in self.pages if p.strip())

    @property
    def visual_pages(self) -> int:
        return self.unread_pages

    @property
    def coverage_complete(self) -> bool:
        """关键范围是否完整覆盖。未读/失败页存在时语义结论必须 unknown。"""
        return self.unread_pages == 0 and self.failed_pages == 0


# ---------------------------------------------------------------- OCR 引擎（懒加载单例）
_ocr_engine = None
_ocr_unavailable = False


def _get_ocr():
    """懒加载 RapidOCR；引擎缺失时返回 None（调用方按未读处理，不崩溃）。"""
    global _ocr_engine, _ocr_unavailable
    if _ocr_engine is not None or _ocr_unavailable:
        return _ocr_engine
    try:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    except Exception:
        _ocr_unavailable = True
    return _ocr_engine


def ocr_available() -> bool:
    """OCR 引擎是否可用（供健康检查/冒烟脚本探测）。"""
    return _get_ocr() is not None


def ocr_image(path: Path) -> tuple[str, str]:
    """对单张图片 OCR，返回 (文本, 错误信息)。不抛异常。"""
    engine = _get_ocr()
    if engine is None:
        return "", "OCR 引擎不可用"
    try:
        result, _elapse = engine(str(path))
    except Exception as e:
        return "", f"OCR 执行失败: {e}"
    if not result:
        return "", "OCR 未识别到文字"
    text = chr(10).join(str(item[1]) for item in result
                        if len(item) > 1 and item[1])
    if not text.strip():
        return "", "OCR 未识别到文字"
    return text, ""


def _ocr_rendered_page(page, page_no: int, tmp_dir: Path) -> tuple[str, str]:
    """把 PDF 页渲染为图片后 OCR，返回 (文本, 错误)。渲染临时文件用后即删。"""
    img_path = tmp_dir / f"ocr_page_{page_no}.png"
    try:
        pix = page.get_pixmap(dpi=150)
        pix.save(str(img_path))
    except Exception as e:
        return "", f"页面渲染失败: {e}"
    try:
        return ocr_image(img_path)
    finally:
        try:
            img_path.unlink()
        except OSError:
            pass


# ---------------------------------------------------------------- 各格式解析
def _parse_pdf(path: Path) -> ParseResult:
    """PDF：有文字层直接提取；无文字层页按需 OCR（预算内）。"""
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ParseError(f"PDF 解析组件缺失: {e}") from e
    try:
        doc = fitz.open(str(path))
    except Exception as e:
        raise ParseError(f"PDF 损坏或无法解析: {e}") from e
    try:
        if doc.needs_pass and not doc.authenticate(""):
            raise ParseError("PDF 已加密，无法读取")
        pages: list[str] = []
        ocr_pages = 0
        unread = 0
        deadline = time.monotonic() + settings.ocr_time_budget_seconds
        ocr_budget = settings.max_ocr_pages_per_file
        for i, page in enumerate(doc):
            try:
                text = page.get_text("text") or ""
            except Exception:
                text = ""
            if text.strip():
                pages.append(text)
                continue
            # 无文字层：扫描件或图片页，按需 OCR（受页数与时限预算约束）
            if ocr_budget <= 0 or time.monotonic() > deadline:
                unread += 1
                continue
            ocr_budget -= 1
            text, _err = _ocr_rendered_page(page, i + 1, path.parent)
            if text:
                pages.append(text)
                ocr_pages += 1
            else:
                unread += 1  # OCR 也没读出：未读页，语义结论必须 unknown
        method = "text"
        if pages and ocr_pages == len(pages):
            method = "ocr"
        elif ocr_pages:
            method = "mixed"
        elif not pages:
            method = "none"
        return ParseResult(pages=pages, method=method,
                           unread_pages=unread, ocr_pages=ocr_pages)
    finally:
        doc.close()


def _parse_docx(path: Path) -> ParseResult:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise ParseError(f"DOCX 无法解析: {e}") from e
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    if not parts:
        return ParseResult(method="none", failed_pages=1, error="DOCX 无可用文本")
    return ParseResult(pages=[chr(10).join(parts)], method="text")


def _parse_text(path: Path) -> ParseResult:
    for enc in ("utf-8", "gb18030"):
        try:
            return ParseResult(pages=[path.read_text(encoding=enc)], method="text")
        except UnicodeDecodeError:
            continue
    raise ParseError("文本编码无法识别")


def _parse_image(path: Path) -> ParseResult:
    """图片：整体作为一页 OCR。识别不了记未读（V07：不写成未填写）。"""
    text, err = ocr_image(path)
    if text:
        return ParseResult(pages=[text], method="ocr", ocr_pages=1)
    return ParseResult(method="visual", unread_pages=1,
                       error=err or "图片无法 OCR，标记未读")


# 文件头魔数（按真实内容判断类型，不信任扩展名）
_MAGIC_PDF = b"%PDF"
_MAGIC_ZIP = b"PK" + bytes([0x03, 0x04])
_MAGIC_JPEG = bytes([0xFF, 0xD8, 0xFF])
_MAGIC_PNG = bytes([0x89]) + b"PNG"
_MAGIC_GIF = b"GIF8"
_MAGIC_RIFF = b"RIFF"
_MAGIC_OLE2 = bytes([0xD0, 0xCF, 0x11, 0xE0])


def parse_file(path: str | Path) -> ParseResult:
    """按真实内容判断类型（不只看扩展名），提取文本。"""
    path = Path(path)
    if not path.exists():
        raise ParseError("文件不存在")
    head = path.read_bytes()[:16]
    if head.startswith(_MAGIC_PDF):
        result = _parse_pdf(path)
        result.detected = "pdf"
        return result
    if head.startswith(_MAGIC_ZIP):
        # zip 容器：DOCX 或伪装格式
        try:
            result = _parse_docx(path)
            result.detected = "docx"
            return result
        except ParseError:
            raise ParseError("ZIP 容器不是受支持的 DOCX，按类型伪装处理")
    if head.startswith((_MAGIC_JPEG, _MAGIC_PNG, _MAGIC_GIF, _MAGIC_RIFF)):
        result = _parse_image(path)
        result.detected = "image"
        return result
    if head.startswith(_MAGIC_OLE2):
        raise ParseError("旧版 Office 文档（OLE2）不受支持，请另存为 .docx 后提交")
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        result = _parse_text(path)
        result.detected = "text"
        return result
    raise ParseError(f"不支持的文件类型: {suffix or '未知'}")
